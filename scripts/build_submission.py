"""Build the cover letter as .docx and assemble a clean submission folder."""
from pathlib import Path
import re, shutil
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
M = ROOT / 'manuscript'
OUT = ROOT / 'submission_package'

text = (M / 'EDITORIAL_MATERIAL.md').read_text(encoding='utf-8')
letter = text.split('## Cover letter')[1].split('\n## ')[0].strip()
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
rf = st.element.get_or_add_rPr().find(qn('w:rFonts'))
for a in ['w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs']:
    rf.set(qn(a), 'Times New Roman')
st.paragraph_format.space_after = Pt(8)
st.font.size = Pt(11.5)
for block in letter.split('\n\n'):
    p = doc.add_paragraph()
    lines = block.split('\n')
    for k, line in enumerate(lines):
        for part in re.split(r'(\*[^*]+\*)', line):
            if part.startswith('*') and part.endswith('*') and len(part) > 2:
                p.add_run(part[1:-1]).italic = True
            elif part:
                p.add_run(part)
        if k < len(lines) - 1:
            p.add_run().add_break()
    if len(lines) > 1:
        p.paragraph_format.space_after = Pt(0)
doc.core_properties.author = 'G. Blake Pierpoint'
doc.save(M / 'cover_letter.docx')

if OUT.exists():
    shutil.rmtree(OUT)
(OUT / 'figures').mkdir(parents=True)
for name in ['manuscript.docx', 'manuscript.pdf', 'supplement.docx', 'supplement.pdf', 'cover_letter.docx']:
    shutil.copy2(M / name, OUT / name)
for name in ['figure1_annual_prevalence.tiff', 'figure2_adjusted_associations.tiff', 'graphic_abstract.jpg']:
    shutil.copy2(M / 'figures' / name, OUT / 'figures' / name)
(OUT / 'README.txt').write_text(
    'Stroke submission files: Cost-Related Barriers to Care and Medications Among US Stroke Survivors, 2019-2025\n\n'
    'manuscript.docx / manuscript.pdf  - main manuscript (title page, abstract, text, references, tables, figures with legends)\n'
    'supplement.docx / supplement.pdf  - Supplemental Material (methods, Figure S1, Tables S1-S12 incl. STROBE checklist)\n'
    'cover_letter.docx                  - cover letter to the editor\n'
    'figures/                           - Figures 1 and 2 (300-dpi TIFF) and the graphic abstract (JPG)\n', encoding='utf-8')
print('submission package:', sorted(str(p.relative_to(OUT)) for p in OUT.rglob('*') if p.is_file()))
