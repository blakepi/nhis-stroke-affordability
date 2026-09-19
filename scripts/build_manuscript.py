"""Build the submission-formatted manuscript, supplement and tables from verified outputs."""
from pathlib import Path
import csv
import json
import re
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[1]
DEST=ROOT/'manuscript'; TABLES=DEST/'tables'; TABLES.mkdir(exist_ok=True)
OUT=ROOT/'outputs/full'
co=pd.read_csv(OUT/'model_coefficients.csv')
components=pd.read_csv(OUT/'component_prevalence.csv').set_index('outcome')
sg=pd.read_csv(OUT/'subgroup_prevalence.csv').set_index(['variable','level'])
adj=pd.read_csv(OUT/'adjusted_annual_prevalence.csv').set_index('year')
log=pd.read_csv(OUT/'logistic_sensitivity.csv').set_index(['variable','level'])
refs=json.loads((DEST/'references.json').read_text(encoding='utf-8'))

def pct(r): return f"{100*r.estimate:.1f}% (95% CI, {100*r.lower:.1f}%–{100*r.upper:.1f}%)"
def ratio(r): return f"{r.PR:.2f} (95% CI, {r.lower:.2f}–{r.upper:.2f})"
def pr(model,term): return co[(co.model==model)&(co.term==term)].iloc[0]
def shortpr(r): return f"{r.PR:.2f} ({r.lower:.2f}–{r.upper:.2f})"
def shortpct(r): return f"{100*r.estimate:.1f} ({100*r.lower:.1f}–{100*r.upper:.1f})"
def table(name,caption,headers,rows,note=''):
    with (TABLES/f'{name}.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f);w.writerow(headers);w.writerows(rows)
    clean=lambda x:str(x).replace('|','/').replace('\n',' ')
    text=f'### {caption}\n\n'+'| '+' | '.join(headers)+' |\n'+'| '+' | '.join(['---']*len(headers))+' |\n'
    text+=''.join('| '+' | '.join(map(clean,row))+' |\n' for row in rows)
    return text+'\n'+note+'\n\n'

labels={
 'young':{'65+':'Age 65 years or older','18-64':'Age 18–64 years'},
 'sex':{'1':'Male','2':'Female'},
 'race':{'1':'Hispanic, any race','2':'Non-Hispanic White only','3':'Non-Hispanic Black only','4':'Non-Hispanic Asian only','5':'Non-Hispanic AIAN only','6':'Non-Hispanic AIAN with another race','7':'Other single or multiple races'},
 'region':{'1':'Northeast','2':'Midwest','3':'South','4':'West'},
 'hypertension':{'No':'No hypertension','Yes':'Hypertension'},
 'diabetes':{'No':'No diabetes','Yes':'Diabetes'},
 'coronary':{'No':'No coronary heart disease','Yes':'Coronary heart disease'},
 'income_group':{'400%+':'Income ≥400% FPL','200-399%':'Income 200–399% FPL','100-199%':'Income 100–199% FPL','<100%':'Income <100% FPL'},
 'smoking':{'Never':'Never smoked','Former':'Former smoker','Current':'Current smoker'},
 'education':{'Bachelor or higher':"Bachelor's degree or higher"},
 'insurance':{'Private':'Private, with or without public coverage'},
 'fairpoor':{'Good or better':'Good or better self-rated health','Fair or poor':'Fair or poor self-rated health'}}
varlabels={'young':'Age','sex':'Sex','race':'Race and Hispanic origin','region':'Region','education':'Education','income_group':'Income','insurance':'Insurance','hypertension':'Hypertension','diabetes':'Diabetes','coronary':'Coronary heart disease','smoking':'Smoking','fairpoor':'Self-rated health','disability':'Disability'}
def level_label(v,l): return (varlabels[v]+' unknown') if l in ['Unknown','Missing'] else labels.get(v,{}).get(str(l),str(l))
t1=pd.read_csv(OUT/'table1_characteristics.csv',dtype={'level':str}).set_index(['group','variable','level'])
groups=['Overall','No barrier','Any barrier']
def characteristic_rows(variables):
    rows=[]
    for v in variables:
        subset=t1.loc['Overall',v]
        for lv,r in subset.iterrows():
            if r.n_mean==0:continue
            cells=[]
            for g in groups:
                z=t1.loc[g,v,lv];n=f'{z.n_mean:,.1f}' if v=='income_group' else f'{int(z.n_mean):,}'
                cells.append(f'{n} ({z.weighted_pct:.1f})')
            rows.append([level_label(v,lv),*cells])
    return rows
t1note='Values are unweighted number (survey-weighted column percentage). Income counts are averaged across the ten NCHS imputations and may be fractional. Unknown values are retained in denominators; percentages within a variable may not sum to 100 because of rounding. FPL indicates federal poverty level; NHIS, National Health Interview Survey. Additional characteristics are given in Table S3.'
main_tables=table('table1','Table 1. Characteristics of US stroke survivors in the pooled 2019–2025 NHIS sample, overall and by cost-related barrier status',['Characteristic','Overall (N=7,104)','No barrier (n=5,920)','Any barrier (n=1,184)'],characteristic_rows(['young','sex','income_group','insurance','disability','fairpoor','smoking']),t1note)
ann=pd.read_csv(ROOT/'outputs/annual_prevalence.csv');ann=ann[(ann.outcome=='any_barrier')&(ann.group=='All survivors')]
rows=[[int(r.year),f'{int(r.n):,}',int(r.events),shortpct(r),shortpct(adj.loc[int(r.year)])] for _,r in ann.iterrows()]
main_tables+=table('table2','Table 2. Observed and standardized annual prevalence of any cost-related barrier among US stroke survivors, 2019–2025',['Year','No. with observed outcome','No. with barrier','Observed prevalence, % (95% CI)','Standardized prevalence, % (95% CI)'],rows,'Observed estimates use each year\'s full sample and annual weights. Standardized estimates come from the fully adjusted logistic model fitted in the pooled complete-case sample with 2020 partial-sample weights and ten income imputations, standardized to the pooled covariate distribution. Confidence intervals account for the complex survey design.')
selected=[('young18-64','Age 18–64 vs ≥65 years'),('sex2','Female vs male'),('educationSome college or associate',"Some college/associate vs bachelor's or higher"),('educationHigh school or GED',"High school/GED vs bachelor's or higher"),('educationLess than high school',"Less than high school vs bachelor's or higher"),('income_group200-399%','Income 200–399% vs ≥400% FPL'),('income_group100-199%','Income 100–199% vs ≥400% FPL'),('income_group<100%','Income <100% vs ≥400% FPL'),('insurancePublic without private','Public without private vs private'),('insuranceMilitary only','Military only vs private'),('insuranceUninsured','Uninsured vs private'),('hypertensionYes','Hypertension vs none'),('diabetesYes','Diabetes vs none'),('coronaryYes','Coronary heart disease vs none'),('smokingFormer','Former vs never smoking'),('smokingCurrent','Current vs never smoking'),('fairpoorFair or poor','Fair/poor vs good or better health'),('disabilityWith disability','Disability vs none')]
rows=[]
for term,label in selected:
    cells=[]
    for m in ['M1','M2','M3']:
        name='Age_'+m if term=='young18-64' else m
        a=co[(co.model==name)&(co.term==term)]
        cells.append(shortpr(a.iloc[0]) if len(a) else '—')
    rows.append([label,*cells])
main_tables+=table('table3','Table 3. Prevalence ratios for any cost-related barrier among US stroke survivors under sequential adjustment',['Characteristic','Model 1, PR (95% CI)','Model 2, PR (95% CI)','Model 3, PR (95% CI)'],rows,'Survey-weighted modified Poisson regression with ten income imputations; all models include the same 6,777 survivors (1,117 with a barrier). Model 1 adjusts for survey year, age (spline), sex, race and Hispanic origin, and region. Model 2 adds education, family income, and insurance. Model 3 adds hypertension, diabetes, coronary heart disease, smoking, self-rated health, and disability. The age-group estimate is from a parallel model that replaces the age spline with an age-group indicator. Private coverage (with or without public coverage) is the reference for insurance. A dash indicates that the variable was not in the model. FPL indicates federal poverty level; GED, General Educational Development; PR, prevalence ratio. Complete Model 3 estimates are given in Table S5.')

supp=''
flow=pd.read_csv(ROOT/'outputs/cohort_flow.csv');poolflow=pd.read_csv(ROOT/'outputs/pooled_cohort_flow.csv').set_index('year')
rows=[[int(r.year),int(r.stroke_eligible),int(r.primary_observed),int(poolflow.loc[r.year,'stroke']),int(poolflow.loc[r.year,'primary_observed']),int(poolflow.loc[r.year,'primary_events'])] for _,r in flow.iterrows()]
supp+=table('tableS1','Table S1. Annual and pooled sample counts by survey year',['Year','Annual stroke','Annual outcome observed','Pooled stroke','Pooled outcome observed','Pooled events'],rows,'Annual stroke records total 7,555; pooled stroke records total 7,181. The difference comes from excluding 2020 followback records in pooled analyses. These are unweighted counts.')
missing=pd.read_csv(OUT/'covariate_missingness.csv')
supp+=table('tableS2','Table S2. Missing nonincome covariates in the pooled stroke-survivor sample',['Variable','Missing n','Missing %'],[[r.variable,int(r.missing),f'{r.pct_missing:.2f}'] for _,r in missing.iterrows()],'Denominator is 7,181 stroke records. Missingness overlaps across variables; the counts must not be summed to derive complete cases. Among records with observed primary outcomes, 327 had at least one missing adjustment variable.')
supp+=table('tableS3','Table S3. Complete characteristics of the pooled sample, overall and by cost-related barrier status',['Characteristic','Overall (N=7,104)','No barrier (n=5,920)','Any barrier (n=1,184)'],characteristic_rows(list(varlabels)),t1note.replace(' Additional characteristics are given in Table S3.','')+' AIAN, American Indian or Alaska Native. The race/Hispanic-origin categories follow the published public-use recode.')
out_labels={'any_barrier':'Any affordability barrier','universal_barrier':'Three universal questions','delayed_care':'Delayed medical care','forgone_care':'Forgone medical care','forgone_rx':'Forgone needed prescriptions','rx_underuse':'Medication underuse among users'}
supp+=table('tableS4','Table S4. Pooled prevalence of the composite outcome and each component barrier, 2019–2025',['Outcome','No. with observed outcome','No. with barrier','Prevalence, % (95% CI)'],[[out_labels[o],int(r.n),int(r.events),shortpct(r)] for o,r in components.iterrows()],'Components overlap. Medication-underuse estimates use prescription users only. Each prevalence denominator includes all stroke records with a classifiable corresponding outcome; it is not restricted to complete model covariates. All displayed estimates pass the implemented precision screen.')
termmap=dict(selected)
termmap.update({f'yearf{y}':f'Year {y} vs 2019' for y in range(2020,2026)})
termmap.update({f'race{i}':f"{labels['race'][str(i)]} vs Hispanic" for i in range(2,8)})
termmap.update({f'region{i}':f"{labels['region'][str(i)]} vs Northeast" for i in range(2,5)})
rows=[]
for _,r in co[co.model=='M3'].iterrows():
    if r.term=='(Intercept)' or r.term.startswith('splines::'):continue
    rows.append([termmap.get(r.term,r.term),shortpr(r),'<0.001' if r.p<.001 else f'{r.p:.3f}'])
supp+=table('tableS5','Table S5. Fully adjusted (Model 3) prevalence ratios for all covariates',['Characteristic','PR (95% CI)','P value'],rows,'The model includes an age spline with 3 degrees of freedom. Spline basis coefficients and the intercept are retained in model_coefficients.csv; they do not individually represent clinical group contrasts. No multiplicity correction was applied. Sparse race and military-coverage contrasts require caution.')
senslabels={'No2020':'Exclude 2020','WorkingAge':'Age 18–64 only','NoDisability':'No disability','Universal':'Universal composite','Delayed':'Delayed medical care','ForgoneCare':'Forgone medical care','ForgoneRx':'Forgone prescriptions','RxUnderuse':'Medication underuse','MissingCategory':'Unknown-category diagnostic'}
diag=pd.read_csv(OUT/'model_diagnostics.csv');rows=[]
for m,label in senslabels.items():
    dm=diag[(diag.model==m)&(diag.imputation==1)].iloc[0]
    cells=[]
    for term in ['income_group100-199%','insuranceUninsured','disabilityWith disability']:
        c=co[(co.model==m)&(co.term==term)]
        cells.append(shortpr(c.iloc[0]) if len(c) else 'Not applicable')
    rows.append([label,f'{int(dm.n):,}',*cells])
supp+=table('tableS6','Table S6. Selected prevalence ratios in sensitivity analyses',['Analysis','No.','Income 100–199% vs ≥400% FPL','Uninsured vs private','Disability vs none'],rows,'Values are PRs (95% CIs) from the corresponding fully adjusted model. No-disability models omit the disability term. The complete coefficient sets, including all income categories, are supplied as CSV. Working-age disability estimates are compatible with no association. Outcome-specific model denominators differ from Table S4 because they require complete adjustment variables.')
rows=[]
for (v,lv),r in log.iterrows():rows.append([level_label(v,lv)+' vs '+level_label(v,r.reference),shortpr(r),f'{100*r.p_reference:.1f}',f'{100*r.p_level:.1f}'])
supp+=table('tableS7','Table S7. Standardized prevalence ratios from bounded logistic models',['Contrast','Marginal PR (95% CI)','Reference %','Comparison %'],rows,'Logistic predictions remain between zero and one. Marginal PRs are pooled on the log scale; displayed prevalences are arithmetic means of imputation-specific margins, so their displayed ratio may differ slightly from the pooled PR. The age contrast uses the age-group model. These standardized estimates have a different estimand from conditional modified Poisson PRs.')
rows=[]
for (v,lv),r in sg.iterrows():
    if r.reliability!='adequate':continue
    n=f'{r.n_mean:.1f}' if v=='income_group' else f'{int(r.n_mean):,}'
    rows.append([level_label(v,lv),n,shortpct(r)])
supp+=table('tableS8','Table S8. Unadjusted pooled prevalence of any cost-related barrier in principal subgroups',['Subgroup','No. with observed outcome','Prevalence, % (95% CI)'],rows,'Income counts are averages across ten imputations. The unknown-insurance subgroup (29 records) failed the precision screen, and its prevalence is not displayed. These are unadjusted subgroup estimates.')
rows=[]
for m,g in diag.groupby('model',sort=False):rows.append([m,int(g.n.iloc[0]),int(g.events.iloc[0]),f'{g.above_one.min()}–{g.above_one.max()}',f'{g.max_fitted.max():.3f}','All converged'])
supp+=table('tableS9','Table S9. Convergence and fitted-value diagnostics for the modified Poisson models',['Model','n','Events','Fitted >1 per imputation','Largest fitted value','Convergence'],rows,'Ten fits per specification. Fitted log-link values above one are not interpreted as individual probabilities. The bounded logistic sensitivity supports the major associations; it does not establish perfect specification. Model labels correspond to the accompanying script and coefficient CSV.')

outcome_definitions=table('definitions_outcomes','Table S10. Outcome definitions and NHIS source variables',['Measure','NHIS variables','Universe'],[['Stroke history','STREV_A','All sample adults'],['Delayed medical care','MEDDL12M_A','All sample adults'],['Forgone medical care','MEDNG12M_A','All sample adults'],['Forgone needed prescriptions','RXDG12M_A','All sample adults'],['Prescription use','RX12M_A','All sample adults'],['Dose skipping, dose reduction, delayed filling','RXSK12M_A; RXLS12M_A; RXDL12M_A','RX12M_A=1'],['Any affordability barrier','Any applicable positive among six items','Known applicability and response rules described below']])
covariate_definitions=table('definitions_covariates','Table S11. Covariate definitions and NHIS source variables',['Domain','Variables','Analysis categories'],[['Age','AGEP_A','18–85; 85 is top-coded; natural spline or 18–64/65+'],['Sex','SEX_A','Male; female'],['Race/Hispanic origin','HISPALLP_A','Seven published categories'],['Region','REGION','Northeast; Midwest; South; West'],['Education','EDUC_A; EDUCP_A','Four harmonized attainment groups'],['Income','RATCAT_A; IMPNUM or IMPNUM_A','<100%; 100–199%; 200–399%; ≥400% FPL'],['Insurance','NOTCOV_A and seven coverage indicators','Private; public without private; military only; uninsured'],['Clinical history','HYPEV_A; DIBEV_A; CHDEV_A','Yes/no hypertension, diabetes, coronary disease'],['Smoking','SMKCIGST_A','Never; former; current; indeterminate missing'],['Self-rated health','PHSTAT_A','Fair/poor vs good/very good/excellent'],['Disability','DISAB3_A','Published Washington Group disability recode']])
strobe_rows=[['1','Title and Abstract','Design, population, period and principal quantitative findings.'],['2–3','Introduction','Background, prior studies and objectives stated.'],['4–5','Methods / Study design','Repeated cross-sectional design, public NHIS setting and 2019–2025 period.'],['6','Methods / Population; Figure S1','Eligibility, stroke definition, age exclusion, annual vs pooled selection.'],['7–8','Methods / Outcomes and covariates','Definitions, universes and cross-year source verification; crosswalk files included.'],['9','Discussion / Limitations','Self-report, nonresponse, institutional exclusion, confounding and missingness addressed.'],['10','Methods / Population','All eligible records in the selected years; no formal power-based sample selection.'],['11','Methods / Covariates','Age spline and categorical income/education definitions given.'],['12','Methods / Analysis; Supplementary Methods','Survey weighting, MI, missingness, interaction and sensitivity procedures.'],['13','Results / Sample; Table S1; Figure S1','Counts and reasons for exclusions; different annual and pooled paths.'],['14','Table 1; Tables S2–S3','Characteristics and missingness; income counts averaged across imputations.'],['15','Table 2; Tables S4 and S8','Outcome counts and annual, component and subgroup estimates.'],['16','Table 3; Table S5; Results','Adjusted PRs, covariate sets, precision and absolute annual difference.'],['17','Sensitivity Results; Tables S6–S9','Sensitivity estimates and model diagnostics.'],['18–21','Discussion and Conclusions','Key results, limitations, cautious interpretation and generalizability.'],['22','Declarations','No external funding is reported.']]
strobe=table('STROBE_checklist','Table S12. STROBE checklist for cross-sectional studies',['Item','Location in manuscript','Content'],strobe_rows,'')

JABBR={'Circulation':'Circulation','Stroke':'Stroke','Archives of Neurology':'Arch Neurol','Annals of Neurology':'Ann Neurol','Journal of Health Care for the Poor and Underserved':'J Health Care Poor Underserved','Journal of Racial and Ethnic Health Disparities':'J Racial Ethn Health Disparities','Journal of Managed Care and Specialty Pharmacy':'J Manag Care Spec Pharm','Journal of Statistical Software':'J Stat Softw','American Journal of Epidemiology':'Am J Epidemiol','Biometrics':'Biometrics','Vital and Health Statistics Series 2':'Vital Health Stat 2','PLoS Medicine':'PLoS Med'}
def cite(r):
    authors=', '.join(r['authors']).rstrip('.')
    if r.get('journal'):
        j=JABBR.get(r['journal'],r['journal'])
        if r.get('online'):
            tail=f"{j}. Published online {r['online']}. doi:{r['doi']}"
        else:
            tail=f"{j}. {r['year']};{r.get('volume','')}"+(f"({r['issue']})" if r.get('issue') else '')+(f":{r['pages']}" if r.get('pages') else '')+'.'
            if r.get('doi'):tail+=f" doi:{r['doi']}"
        return f"{r['id']}. {authors}. {r['title']}. {tail}"
    return f"{r['id']}. {authors}. {r['title']}. National Center for Health Statistics; {r['year']}. Accessed September 18, 2026. {r['url']}"
references='\n\n'.join(cite(r) for r in refs)
difference=pd.read_csv(OUT/'adjusted_year_difference.csv').iloc[0];tests=pd.read_csv(OUT/'joint_tests.csv')
tokens={'pooled':pct(components.loc['any_barrier']),'young':pct(sg.loc['young','18-64']),'older':pct(sg.loc['young','65+']),
 'agepr':ratio(pr('Age_M3','young18-64')),'incomepr':ratio(pr('M3','income_group100-199%')),'income200pr':ratio(pr('M3','income_group200-399%')),'incomelowpr':ratio(pr('M3','income_group<100%')),
 'uninsuredpr':ratio(pr('M3','insuranceUninsured')),'publicpr':ratio(pr('M3','insurancePublic without private')),'disabilitypr':ratio(pr('M3','disabilityWith disability')),
 'adj2019':pct(adj.loc[2019]),'adj2025':pct(adj.loc[2025]),'difference':f'{100*difference.difference:.1f} percentage points (95% CI, {100*difference.lower:.1f} to {100*difference.upper:.1f})',
 'yearp':f'{tests.iloc[0].p:.3f}','interactionp':f'{tests.iloc[1].p:.3f}',
 'delayed':pct(components.loc['delayed_care']),'forgonecare':pct(components.loc['forgone_care']),'forgonerx':pct(components.loc['forgone_rx']),'underuse':pct(components.loc['rx_underuse']),'universal':pct(components.loc['universal_barrier']),
 'uninsuredprev':pct(sg.loc['insurance','Uninsured']),'privateprev':pct(sg.loc['insurance','Private']),'publicprev':pct(sg.loc['insurance','Public without private']),'disabilityprev':pct(sg.loc['disability','With disability']),'nodisabilityprev':pct(sg.loc['disability','Without disability']),
 'logage':ratio(log.loc['young','18-64']),'logincome':ratio(log.loc['income_group','100-199%']),'loguninsured':ratio(log.loc['insurance','Uninsured']),'logdisability':ratio(log.loc['disability','With disability']),
 'references':references,'main_tables':main_tables,'outcome_definitions':outcome_definitions,'covariate_definitions':covariate_definitions,'supplement_tables':supp,'strobe':strobe}

def inline_runs(p,text):
    runs=[]
    pattern=r'(\*\*.*?\*\*|\^[^^ ]+\^)'
    for part in re.split(pattern,text):
        if not part:continue
        if part.startswith('**') and part.endswith('**'):r=p.add_run(part[2:-2]);r.bold=True
        elif len(part)>2 and part.startswith('^') and part.endswith('^'):r=p.add_run(part[1:-1]);r.font.superscript=True
        else:r=p.add_run(part)
        runs.append(r)
    return runs
def inline(p,text):return inline_runs(p,text)

def add_page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def set_line_numbers(section):
    sectPr=section._sectPr
    ln=OxmlElement('w:lnNumType');ln.set(qn('w:countBy'),'1');ln.set(qn('w:restart'),'continuous');ln.set(qn('w:distance'),'360')
    sectPr.append(ln)

def write_docx(markdown,path,body_spacing=2.0,line_numbers=False,figure_width=6.5):
    doc=Document();s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
        rpr=style.element.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
        if rf is None:rf=OxmlElement('w:rFonts');rpr.append(rf)
        for attr in ['w:ascii','w:hAnsi','w:eastAsia','w:cs']:rf.set(qn(attr),'Times New Roman')
    normal=doc.styles['Normal'];normal.font.size=Pt(12);normal.paragraph_format.line_spacing=body_spacing;normal.paragraph_format.space_after=Pt(0)
    title=doc.styles['Title'];title.font.size=Pt(14);title.font.bold=True;title.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;title.paragraph_format.space_after=Pt(12);title.paragraph_format.line_spacing=1.15
    tpr=title.element.get_or_add_pPr();bdr=tpr.find(qn('w:pBdr'))
    if bdr is not None:tpr.remove(bdr)
    for name,size,italic in [('Heading 1',12,False),('Heading 2',12,True),('Heading 3',12,True)]:
        h=doc.styles[name];h.font.size=Pt(size);h.font.bold=True;h.font.italic=italic;h.paragraph_format.space_before=Pt(12);h.paragraph_format.space_after=Pt(0);h.paragraph_format.keep_with_next=True;h.paragraph_format.line_spacing=body_spacing
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    if line_numbers:set_line_numbers(s)
    lines=markdown.splitlines();i=0;title_page=True
    break_before={'## Abstract','## Introduction','## References','## Tables','## Figure Legends','## Supplementary Tables','## Supplementary Figure','## STROBE Checklist'}
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line in break_before:add_page_break(doc);title_page=False
        if line.startswith('| '):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',x) for x in row):rows.append(row)
                i+=1
            count=len(rows[0]);tab=doc.add_table(rows=1,cols=count);tab.alignment=WD_TABLE_ALIGNMENT.CENTER;tab.autofit=False
            first=2.7 if count==4 else 2.3 if count<=3 else 1.8
            widths=[first]+[(6.5-first)/(count-1)]*(count-1)
            for j,w in enumerate(widths):tab.columns[j].width=Inches(w)
            borders=OxmlElement('w:tblBorders')
            for edge,sz in [('top','8'),('bottom','8'),('insideH','4')]:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),sz);e.set(qn('w:color'),'000000' if sz=='8' else 'BFBFBF');borders.append(e)
            for edge in ['left','right','insideV']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'nil');borders.append(e)
            tab._tbl.tblPr.append(borders)
            for k,values in enumerate(rows):
                row=tab.rows[0] if k==0 else tab.add_row()
                if k==0:repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
                for j,value in enumerate(values):
                    cell=row.cells[j];cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p=cell.paragraphs[0];p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                    if j:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    run=p.add_run(value);run.font.size=Pt(9);run.bold=(k==0)
                    if k==0:
                        bb=OxmlElement('w:tcBorders');b=OxmlElement('w:bottom');b.set(qn('w:val'),'single');b.set(qn('w:sz'),'6');b.set(qn('w:color'),'000000');bb.append(b);cell._tc.get_or_add_tcPr().append(bb)
            continue
        if line.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',line);img=DEST/m.group(2)
            width=5.4 if 'cohort' in img.name else figure_width
            add_page_break(doc)
            cap=doc.add_paragraph();cap.paragraph_format.line_spacing=1.15;cap.add_run(m.group(1)).bold=True
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(img),width=Inches(width));i+=1;continue
        if line.startswith('#'):
            n=len(line)-len(line.lstrip('#'));text=line[n:].strip()
            if n==1:
                doc.add_paragraph(text,style='Title')
            elif re.match(r'^Table S?\d+\.',text):
                if re.match(r'^Table [2-9]\.',text):add_page_break(doc)
                p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(4)
                m=re.match(r'^(Table S?\d+\.)\s*(.*)$',text);p.add_run(m.group(1)+' ').bold=True;p.add_run(m.group(2))
            else:
                doc.add_paragraph(text,style='Heading '+str(min(n-1,3)))
            i+=1;continue
        p=doc.add_paragraph()
        j=i-1
        while j>=0 and not lines[j].strip():j-=1
        prev=lines[j].strip() if j>=0 else ''
        is_note=prev.startswith('|')
        if is_note:
            p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(6)
            for r in inline_runs(p,line):r.font.size=Pt(9)
        elif title_page:
            p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(8)
            if i<=3:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            inline(p,line)
        else:
            inline(p,line)
        i+=1
    doc.core_properties.title=markdown.splitlines()[0].lstrip('# ')
    doc.core_properties.subject='Cost-related barriers to care and medications among US stroke survivors, NHIS 2019-2025'
    doc.core_properties.author='G. Blake Pierpoint; Alberto E. Musto'
    doc.save(path)

metadata={}
for source,outname in [('draft_source.md','manuscript'),('supplement_source.md','supplement')]:
    text=(DEST/source).read_text(encoding='utf-8')
    for key,value in tokens.items():text=text.replace('{{'+key+'}}',value)
    if outname=='manuscript':
        abstract_n=len(text.split('## Abstract')[1].split('## Introduction')[0].split())
        body_n=len(text.split('## Introduction')[1].split('## Declarations')[0].split())
        text=text.replace('{{counts}}',f'Word count: abstract, {abstract_n}; main text, {body_n:,}. References: {len(refs)}. Tables: 3. Figures: 2. Supplementary material: Supplementary Methods, Tables S1–S12, Figure S1.')
    assert not re.search(r'\{\{[^}]+\}\}',text),'Unresolved numeric/content token'
    (DEST/f'{outname}.md').write_text(text,encoding='utf-8')
    write_docx(text,DEST/f'{outname}.docx',body_spacing=2.0 if outname=='manuscript' else 1.15,line_numbers=(outname=='manuscript'))
    metadata[outname]={'words_including_tables_and_references':len(text.split()),'tables':text.count('\n| ---'),'images':len(re.findall(r'!\[',text))}
    if outname=='manuscript':
        metadata[outname]['abstract_words']=len(text.split('## Abstract')[1].split('## Introduction')[0].split())
        metadata[outname]['body_words']=len(text.split('## Introduction')[1].split('## Declarations')[0].split())
ris=[]
for r in refs:
    a=['TY  - JOUR' if r.get('journal') else 'TY  - RPRT',f"ID  - {r['id']}",f"TI  - {r['title']}",f"PY  - {r['year']}"]
    a.extend('AU  - '+x for x in r['authors'] if x!='et al.')
    for key,tag in [('journal','JO'),('volume','VL'),('issue','IS'),('pages','SP'),('doi','DO'),('url','UR')]:
        if r.get(key):a.append(f'{tag}  - {r[key]}')
    ris.append('\n'.join(a+['ER  -','']))
(DEST/'references.ris').write_text('\n'.join(ris),encoding='utf-8')
(OUT/'manuscript_metadata.json').write_text(json.dumps(metadata,indent=2),encoding='utf-8')
print(json.dumps(metadata,indent=2))
